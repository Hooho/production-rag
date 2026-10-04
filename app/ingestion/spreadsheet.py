"""Excel（.xlsx）和 CSV 的解析：把每个工作表里的表格转成「列名：值」一行一段的章节，供分块器按行切分。

为什么不用 Unstructured：它把整个工作表转成一段 HTML 表格，表头、行号、多张表的边界都没了，
按字数切开以后后面的分片没有表头，模型不知道「张三 | 3500」是什么意思。

处理规则（每个工作表单独处理，互不影响）：
1. 隐藏的工作表、行、列跳过；合并单元格把值填满整个合并区域（纵向合并的「华东区」每行都带上）。
2. 用整行全空、整列全空把工作表切成几块，每块当一张独立的表。
3. 块里只有一个值的行（或横跨整行的合并单元格）是表格标题；只有标题、说明文字的块，
   紧挨着下一张表的短文字当那张表的标题，其余当说明文字（跟在上一张表后面就是这张表的备注）。
4. 表头：在标题下面的第一行，要求大部分格子有值、全是文字、互不重复；上面一行有合并单元格或重复值、
   下面一行也全是文字时按两层（最多三层）表头拼成「上级 / 下级」。
5. 认不出表头时不猜列名，按 Excel 列字母输出「第 12 行：A=张三；B=3500」，并在识别结果里标出来。
每张表、每个工作表的识别结果记进 table_report，文档详情里逐个显示。
"""

import csv
from datetime import date, datetime, time
import io
from pathlib import Path


SPREADSHEET_EXTENSIONS = {".xlsx", ".csv"}
# 表头只在表格开头这么多行里找；再往下还没找到，多半是没有表头的表。
HEADER_SEARCH_ROWS = 5
MAX_HEADER_DEPTH = 3
# 紧挨着表格、不超过这么长的单独一行文字当表格标题；更长的当说明文字。
TITLE_MAX_CHARS = 60
# 以这些开头、或以句末标点结尾的文字是备注、说明，不当表格标题。
NOTE_PREFIXES = ("注", "备注", "说明", "*", "※")
SENTENCE_ENDINGS = ("。", "；", ";", "！", "？", ".", "!", "?")


def column_letter(index):
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


# 数字转成文字：整数不带小数点，小数去掉末尾的 0，避免 0.30000000000000004 这类浮点误差。
def number_text(value):
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        text = f"{round(value, 6):.6f}".rstrip("0").rstrip(".")
        return text if text not in {"", "-0"} else "0"
    return str(value)


# 单元格的值转成 (显示文字, 类型)。类型用来判断表头：表头全是文字，数据行里常有数字和日期。
def cell_text(value, number_format=""):
    if value is None:
        return None, None
    if isinstance(value, bool):
        return ("是" if value else "否"), "bool"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime("%Y-%m-%d"), "date"
        return value.strftime("%Y-%m-%d %H:%M"), "date"
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d"), "date"
    if isinstance(value, time):
        return value.strftime("%H:%M"), "date"
    if isinstance(value, (int, float)):
        if "%" in (number_format or ""):
            return number_text(value * 100) + "%", "number"
        return number_text(value), "number"
    text = " ".join(str(value).split())
    if not text:
        return None, None
    return text, ("number" if looks_numeric(text) else "text")


# CSV 里全是字符串，「3500」「12.5%」「2024-03-01」也要当成数字、日期，不然没法和表头区分。
def looks_numeric(text):
    stripped = text.replace(",", "").replace("%", "").replace("¥", "").replace("$", "").strip()
    if not stripped:
        return False
    try:
        float(stripped)
        return True
    except ValueError:
        pass
    parts = stripped.replace("/", "-").split("-")
    return len(parts) == 3 and all(part.isdigit() for part in parts)


# 读 xlsx：每个可见工作表一个 {name, rows}，rows 是 [(Excel 行号, [(列号, 文字, 类型), ...])]，只含可见的行和列。
def read_xlsx(path):
    from openpyxl import load_workbook

    # data_only=True 读公式的计算结果（Excel 保存时缓存的值），不读公式本身。
    workbook = load_workbook(str(path), data_only=True)
    sheets = []
    hidden_sheets = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            hidden_sheets.append(sheet.title)
            continue
        # 合并单元格只有左上角有值，先把值填满整个合并区域。
        filled = {}
        for merged in sheet.merged_cells.ranges:
            anchor = sheet.cell(merged.min_row, merged.min_col)
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    filled[(row, col)] = (anchor.value, anchor.number_format)
        hidden_cols = {key for key, dimension in sheet.column_dimensions.items() if dimension.hidden}
        rows = []
        for row_cells in sheet.iter_rows():
            if not row_cells:
                continue
            row_number = row_cells[0].row
            if sheet.row_dimensions[row_number].hidden:
                continue
            cells = []
            for cell in row_cells:
                if column_letter(cell.column - 1) in hidden_cols:
                    continue
                value, number_format = filled.get((row_number, cell.column), (cell.value, cell.number_format))
                text, kind = cell_text(value, number_format)
                if text is not None:
                    cells.append((cell.column - 1, text, kind))
            rows.append((row_number, cells))
        sheets.append({"name": sheet.title, "rows": rows})
    return sheets, hidden_sheets


# 读 csv：编码先试 UTF-8（带不带 BOM 都行），不行再按 GB18030（兼容 GBK、GB2312）；分隔符自动识别。
def read_csv(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row_number, values in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        cells = []
        for col, value in enumerate(values):
            cell, kind = cell_text(value)
            if cell is not None:
                cells.append((col, cell, kind))
        rows.append((row_number, cells))
    return [{"name": None, "rows": rows}], []


# 用整行全空、整列全空把工作表切成几块：先按空行分段，每段里再按空列分开。
def split_blocks(rows):
    bands = []
    current = []
    for row_number, cells in rows:
        if cells:
            current.append((row_number, cells))
        elif current:
            bands.append(current)
            current = []
    if current:
        bands.append(current)
    blocks = []
    for band in bands:
        used = sorted({col for _, cells in band for col, _, _ in cells})
        groups = []
        for col in used:
            if groups and col == groups[-1][-1] + 1:
                groups[-1].append(col)
            else:
                groups.append([col])
        for group in groups:
            columns = set(group)
            block_rows = []
            for row_number, cells in band:
                picked = [cell for cell in cells if cell[0] in columns]
                if picked:
                    block_rows.append((row_number, picked))
            if block_rows:
                blocks.append({"columns": group, "rows": block_rows})
    return blocks


def distinct_values(cells):
    seen = []
    for _, text, _ in cells:
        if text not in seen:
            seen.append(text)
    return seen


# 只有一个值的行：标题、说明文字，或者横跨整行的合并单元格（填满后每格都是同一个值）。
def is_single_value(cells):
    return len(distinct_values(cells)) == 1


# 一行能不能当表头的一层：大部分格子有值，全是文字，每格不太长。
def header_like(cells, width):
    if not cells or len(cells) < max(2, (width + 1) // 2):
        return False
    if any(kind != "text" for _, _, kind in cells):
        return False
    return all(len(text) <= 30 for _, text, _ in cells)


# 把几层表头拼成每列的列名：上层空着的格子沿用左边的值（没合并、只是留空的分组表头），
# 上下两层一样（纵向合并）只写一次。
def combine_header(levels, columns):
    names = {}
    filled_levels = []
    for depth, cells in enumerate(levels):
        values = {col: text for col, text, _ in cells}
        if depth < len(levels) - 1:
            last = None
            for col in columns:
                if col in values:
                    last = values[col]
                elif last is not None and any(col in {c for c, _, _ in lower} for lower in levels[depth + 1:]):
                    values[col] = last
        filled_levels.append(values)
    for col in columns:
        parts = []
        for values in filled_levels:
            text = values.get(col)
            if text and text not in parts:
                parts.append(text)
        if parts:
            names[col] = " / ".join(parts)
    return names


# 在块的开头找表头，返回 (标题行, 表头层数, 列名, 表头的 Excel 行号, 置信度)；找不到时层数为 0。
def detect_header(block):
    rows = block["rows"]
    columns = block["columns"]
    width = len(columns)
    titles = []
    start = 0
    # 开头只有一个值的行是表格标题（多列的表才这样判断；单列的表本来每行就只有一个值）。
    while width > 1 and start < len(rows) - 1 and is_single_value(rows[start][1]):
        titles.append(distinct_values(rows[start][1])[0])
        start += 1
    for offset in range(start, min(start + HEADER_SEARCH_ROWS, len(rows) - 1)):
        if not header_like(rows[offset][1], width):
            continue
        candidates = []
        for depth in range(1, MAX_HEADER_DEPTH + 1):
            if offset + depth >= len(rows):
                break
            levels = [cells for _, cells in rows[offset:offset + depth]]
            # 下面几层也得全是文字才可能是表头。
            if depth > 1 and not all(cells and all(kind == "text" for _, _, kind in cells) for cells in levels[1:]):
                break
            names = combine_header(levels, columns)
            if len(names) >= max(2, (width + 1) // 2) and len(set(names.values())) == len(names):
                candidates.append((len(names), -depth, depth, names))
        if not candidates:
            continue
        # 第一层就把每列区分开时只用一层；第一层有合并（「2024年 2024年 2025年 2025年」）或留空时，
        # 用能给最多列起出不重复列名的层数，一样多时用更少的层。
        _, _, depth, names = max(candidates)
        header_rows = [row_number for row_number, _ in rows[offset:offset + depth]]
        data = rows[offset + depth:]
        # 整张表都是文字时（比如联系人表），第一行是表头还是数据只能推测，标成低置信度。
        has_typed = any(kind != "text" for _, cells in data for _, _, kind in cells)
        confidence = "high" if has_typed else "low"
        return titles, depth, names, header_rows, confidence, offset + depth
    return titles, 0, {}, [], None, start


# 一行数据转成文字：有列名时「列名：值」，没有时「A=值」；只有一个值的行（表中的分组行）直接写值。
def row_text(cells, names):
    if names:
        if len(cells) == 1 and len(names) > 2:
            return cells[0][1]
        parts = []
        for col, text, _ in cells:
            name = names.get(col)
            parts.append(f"{name}：{text}" if name else text)
        return "；".join(parts)
    return "；".join(f"{column_letter(col)}={text}" for col, text, _ in cells)


def block_text(block):
    return "\n".join(" ".join(text for _, text, _ in distinct_cells(cells)) for _, cells in block["rows"])


def distinct_cells(cells):
    result = []
    seen = set()
    for cell in cells:
        if cell[1] not in seen:
            seen.add(cell[1])
            result.append(cell)
    return result


# 只有标题、说明文字的块（每行只有一个值）。
def is_text_block(block):
    return all(is_single_value(cells) for _, cells in block["rows"])


# 把一个工作表转成章节和识别结果。sheet_path 是标题路径的前缀（工作表名；CSV 为空）。
def sheet_sections(sheet, element_offset):
    sheet_name = sheet["name"]
    sheet_path = [sheet_name] if sheet_name else []
    blocks = split_blocks(sheet["rows"])
    sections = []
    tables = []
    pending_title = None
    index = element_offset
    table_blocks = [i for i, block in enumerate(blocks) if not is_text_block(block)]
    for position, block in enumerate(blocks):
        if is_text_block(block):
            text = block_text(block)
            next_is_table = position + 1 < len(blocks) and not is_text_block(blocks[position + 1])
            looks_like_title = not text.startswith(NOTE_PREFIXES) and not text.endswith(SENTENCE_ENDINGS)
            if next_is_table and looks_like_title and len(block["rows"]) <= 2 and len(text) <= TITLE_MAX_CHARS:
                pending_title = text.replace("\n", " ")
                continue
            # 说明文字：跟在表格后面就挂在那张表下面（备注），否则挂在工作表下面。
            path = list(sections[-1]["heading_path"]) if sections and sections[-1].get("table") else list(sheet_path)
            parts = []
            for row_number, cells in block["rows"]:
                parts.append({"text": " ".join(text for _, text, _ in distinct_cells(cells)), "page_number": None,
                    "element_type": "Note", "element_index": index, "author": None,
                    "sheet": sheet_name, "row": row_number})
                index += 1
            sections.append({"heading_path": path, "text": "\n".join(part["text"] for part in parts), "parts": parts})
            continue
        titles, depth, names, header_rows, confidence, data_start = detect_header(block)
        title = " ".join(titles) or pending_title
        pending_title = None
        if not title and len(table_blocks) > 1:
            title = f"表 {table_blocks.index(position) + 1}"
        path = sheet_path + ([title] if title else [])
        parts = []
        single_column = len(block["columns"]) == 1
        for row_number, cells in block["rows"][data_start:]:
            # 单列的表（一列清单）本来就没有列名可言，直接写值；多列又认不出表头时按列字母写，不猜列名。
            if depth:
                text = row_text(cells, names)
            elif single_column:
                text = cells[0][1]
            else:
                text = f"第 {row_number} 行：{row_text(cells, {})}"
            parts.append({"text": text,
                "page_number": None, "element_type": "TableRow", "element_index": index, "author": None,
                "sheet": sheet_name, "row": row_number, "atomic": True})
            index += 1
        data_rows = block["rows"][data_start:]
        tables.append({"title": title, "first_row": block["rows"][0][0], "last_row": block["rows"][-1][0],
            "columns": len(block["columns"]), "rows": len(data_rows),
            "header_rows": header_rows, "header_depth": depth, "confidence": confidence, "single_column": single_column,
            "column_names": [names[col] for col in block["columns"] if col in names][:50]})
        if parts:
            sections.append({"heading_path": path, "text": "\n".join(part["text"] for part in parts), "parts": parts,
                "table": True})
    report = {"sheet": sheet_name, "tables": tables,
        "text_only": not tables and bool(sections), "empty": not sections}
    return sections, report, index


# 解析整个文件，返回 (章节, 识别结果)。识别结果：{"sheets": [...], "hidden_sheets": [...]}。
def spreadsheet_sections(path):
    suffix = Path(path).suffix.lower()
    sheets, hidden = read_xlsx(path) if suffix == ".xlsx" else read_csv(path)
    sections = []
    reports = []
    index = 0
    for sheet in sheets:
        sheet_result, report, index = sheet_sections(sheet, index)
        for section in sheet_result:
            section.pop("table", None)
        sections.extend(sheet_result)
        reports.append(report)
    return sections, {"sheets": reports, "hidden_sheets": hidden}
