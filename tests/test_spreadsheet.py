from datetime import date

from openpyxl import Workbook

from app.ingestion.chunking import chunk_document_records
from app.ingestion.parser import extract_sections_cached, parse_metadata
from app.ingestion.spreadsheet import spreadsheet_sections
from app.models import Models
from test_app import headers, setup  # noqa: F401  复用 API 测试的内存存储夹具


# 每个标题路径下表格行的文字（备注另算）。
def texts(sections):
    result = {}
    for section in sections:
        if section["parts"][0]["element_type"] != "Note":
            result.setdefault(" / ".join(section["heading_path"]), section["text"])
    return result


# 一个工作表里两张表：上方合并单元格是第一张表的标题，「注：」开头的是它的备注，第二张没有标题时叫「表 2」。
def test_multiple_tables_titles_and_notes(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "2024 年华东区报价"
    sheet.merge_cells("A1:D1")
    sheet.append(["产品", "规格", "单价", "起订量"])
    sheet.append(["A100", "500ml", 35, 100])
    sheet.append(["A200", "1L", 60.5, 50])
    sheet.append([])
    sheet.append(["注：以上价格不含税"])
    sheet.append([])
    sheet.append(["区域", "负责人", "电话"])
    sheet.append(["华东", "张三", "13800000000"])
    path = tmp_path / "price.xlsx"
    workbook.save(path)

    sections, report = spreadsheet_sections(path)
    result = texts(sections)
    assert result["报价 / 2024 年华东区报价"].startswith("产品：A100；规格：500ml；单价：35；起订量：100")
    assert "产品：A200；规格：1L；单价：60.5；起订量：50" in result["报价 / 2024 年华东区报价"]
    notes = [section for section in sections if section["parts"][0]["element_type"] == "Note"]
    assert notes[0]["heading_path"] == ["报价", "2024 年华东区报价"] and notes[0]["text"] == "注：以上价格不含税"
    assert result["报价 / 表 2"] == "区域：华东；负责人：张三；电话：13800000000"
    tables = report["sheets"][0]["tables"]
    assert [table["header_rows"] for table in tables] == [[2], [8]]
    assert sections[0]["parts"][0]["row"] == 3 and sections[0]["parts"][0]["atomic"] is True


# 两层表头：上层合并的「2024年」和下层的「一季度」拼成「2024年 / 一季度」；纵向合并的「产品」只写一次。
def test_two_level_header(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销量"
    sheet.append(["产品", "2024年", None, "2025年", None])
    sheet.append([None, "一季度", "二季度", "一季度", "二季度"])
    sheet.merge_cells("B1:C1")
    sheet.merge_cells("D1:E1")
    sheet.merge_cells("A1:A2")
    sheet.append(["A100", 10, 20, 30, 40])
    path = tmp_path / "sales.xlsx"
    workbook.save(path)

    sections, report = spreadsheet_sections(path)
    assert sections[0]["text"] == "产品：A100；2024年 / 一季度：10；2024年 / 二季度：20；2025年 / 一季度：30；2025年 / 二季度：40"
    table = report["sheets"][0]["tables"][0]
    assert table["header_depth"] == 2 and table["header_rows"] == [1, 2]


# 数据区纵向合并的单元格每行都带上；日期按日期写；隐藏的工作表跳过；只有文字的工作表按文字处理；
# 认不出表头的不猜列名，按列字母输出。
def test_merged_rows_hidden_text_and_fallback(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分组"
    sheet.append(["地区", "城市", "销售额", "日期"])
    sheet.append(["华东", "上海", 100, date(2024, 3, 1)])
    sheet.append([None, "杭州", 80, date(2024, 3, 2)])
    sheet.merge_cells("A2:A3")
    numbers = workbook.create_sheet("乱表")
    numbers.append([1, 2, 3])
    numbers.append([4, 5, 6])
    workbook.create_sheet("说明").append(["这是一段说明文字，介绍本表格的用途。"])
    hidden = workbook.create_sheet("隐藏")
    hidden.sheet_state = "hidden"
    hidden.append(["x", "y"])
    path = tmp_path / "mixed.xlsx"
    workbook.save(path)

    sections, report = spreadsheet_sections(path)
    result = texts(sections)
    assert result["分组"] == "地区：华东；城市：上海；销售额：100；日期：2024-03-01\n地区：华东；城市：杭州；销售额：80；日期：2024-03-02"
    assert result["乱表"] == "第 1 行：A=1；B=2；C=3\n第 2 行：A=4；B=5；C=6"
    assert [section["text"] for section in sections if section["heading_path"] == ["说明"]] == ["这是一段说明文字，介绍本表格的用途。"]
    assert "隐藏" not in result and report["hidden_sheets"] == ["隐藏"]
    by_sheet = {item["sheet"]: item for item in report["sheets"]}
    assert by_sheet["乱表"]["tables"][0]["header_depth"] == 0
    assert by_sheet["说明"]["text_only"] is True


# CSV：GBK 编码也能读；全是文字的表（没有数字、日期）表头是推测的，标成低置信度。
def test_csv_encoding_and_confidence(tmp_path):
    path = tmp_path / "staff.csv"
    path.write_bytes("姓名,部门,工资\n张三,销售,3500\n".encode("gb18030"))
    sections, report = spreadsheet_sections(path)
    assert sections[0]["heading_path"] == [] and sections[0]["text"] == "姓名：张三；部门：销售；工资：3500"
    assert report["sheets"][0]["tables"][0]["confidence"] == "high"
    text_only = tmp_path / "contacts.csv"
    text_only.write_text("区域,负责人\n华东,张三\n", encoding="utf-8")
    assert spreadsheet_sections(text_only)[1]["sheets"][0]["tables"][0]["confidence"] == "low"


# 按整行切：一行不会被切成两半，表格分片不重叠，分片记下行号范围。
def test_rows_are_not_split_and_ranges_recorded(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["编号", "名称", "说明"])
    for index in range(60):
        sheet.append([index, f"产品{index}", "这是一段比较长的产品说明文字，用来测试按行切分。" * 2])
    path = tmp_path / "big.xlsx"
    workbook.save(path)
    sections, _ = spreadsheet_sections(path)
    records = chunk_document_records("", sections=sections, chunk_size=800, chunk_overlap=120)
    assert len(records) > 1
    assert records[0]["row_start"] == 2 and records[-1]["row_end"] == 61
    for record in records:
        assert all(line.startswith("编号：") for line in record["content"].split("\n"))
        assert record["effective_overlap"] == 0
    for previous, current in zip(records, records[1:]):
        assert current["row_start"] == previous["row_end"] + 1


# 解析入口：Excel 不走 Unstructured，识别结果记进文档元数据。
def test_parse_entry_records_table_report(tmp_path):
    workbook = Workbook()
    workbook.active.append(["产品", "单价"])
    workbook.active.append(["A100", 35])
    path = tmp_path / "entry.xlsx"
    workbook.save(path)
    sections, stats, cached = extract_sections_cached(path, tmp_path / "cache")
    metadata = parse_metadata(path, sections, stats=stats)
    assert cached is False and metadata["parser"] == "spreadsheet"
    assert metadata["table_report"]["sheets"][0]["tables"][0]["column_names"] == ["产品", "单价"]


# 上传：.xlsx、.csv 可以上传，旧版 .xls 提示另存为 .xlsx；分片接口和检索来源带上行号。
def test_upload_and_chunk_rows(setup, tmp_path, monkeypatch):
    client, store = setup
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path))
    rejected = client.post("/documents/upload", headers=headers(), files={"file": ("old.xls", b"x", "application/vnd.ms-excel")})
    assert rejected.status_code == 415 and ".xlsx" in rejected.json()["detail"]
    accepted = client.post("/documents/upload", headers=headers(), files={"file": ("staff.csv", "姓名,工资\n张三,3500\n".encode(), "text/csv")})
    assert accepted.status_code == 202

    workbook = Workbook()
    workbook.active.title = "报价"
    workbook.active.append(["产品", "单价"])
    workbook.active.append(["A100", 35])
    path = tmp_path / "price.xlsx"
    workbook.save(path)
    sections, _ = spreadsheet_sections(path)
    document_id = "22222222-2222-4222-8222-222222222222"
    from app.mysql.store import documents
    with store.engine.begin() as connection:
        connection.execute(documents.insert().values(id=document_id, owner="alice", title="报价", filename="price.xlsx",
            path=str(path), status="ready:1", error=None, document_metadata={}, created="2026-01-01", updated="2026-01-01"))
    store.ingest("alice", "报价", sections[0]["text"], Models(), document_id, document_id, sections=sections,
        source_format=".xlsx")
    chunk = client.get(f"/documents/{document_id}/chunks", headers=headers()).json()["chunks"][0]
    assert chunk["row_start"] == 2 and chunk["row_end"] == 2
    assert chunk["heading_path"] == ["报价"]
    assert store.vector_row({"id": "x", "owner": "alice", "title": "报价", "text": "", "position": 0, "vector": [],
        "chunk_key": "k"}, document_id, {"heading_path": ["报价"], "row_start": 2, "row_end": 5})["heading"] == "报价 · 第 2–5 行"
